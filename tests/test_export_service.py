import csv
import io

import pytest
from openpyxl import load_workbook

from app.services.export_service import ExportService


@pytest.fixture
def sample_rows():
    return [
        {
            "customer_id": "C1",
            "waiting_time": 0.0,
            "system_time": 2.5,
        },
        {
            "customer_id": "C2",
            "waiting_time": 1.5,
            "system_time": 4.0,
        },
    ]


def test_csv_export_contains_headers_and_data(sample_rows):
    content = ExportService.to_csv(sample_rows)

    rows = list(csv.DictReader(io.StringIO(content)))

    assert len(rows) == 2
    assert rows[0]["customer_id"] == "C1"
    assert float(rows[1]["waiting_time"]) == 1.5


def test_excel_export_contains_headers_and_data(sample_rows):
    content = ExportService.to_excel(sample_rows)

    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook.active

    assert worksheet.cell(1, 1).value == "customer_id"
    assert worksheet.cell(2, 1).value == "C1"
    assert worksheet.cell(3, 2).value == 1.5
    assert worksheet.freeze_panes == "A2"


def test_csv_export_rejects_empty_rows():
    with pytest.raises(ValueError):
        ExportService.to_csv([])


def test_excel_export_rejects_empty_rows():
    with pytest.raises(ValueError):
        ExportService.to_excel([])


def test_export_rejects_invalid_rows():
    with pytest.raises(ValueError):
        ExportService.to_csv(["invalid"])

