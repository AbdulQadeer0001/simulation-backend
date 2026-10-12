import io

import pytest
from openpyxl import load_workbook

from app import create_app


@pytest.fixture
def client():
    app = create_app()
    app.config["TESTING"] = True

    with app.test_client() as test_client:
        yield test_client


@pytest.fixture
def payload():
    return {
        "rows": [
            {
                "customer_id": "C1",
                "waiting_time": 0.0,
                "system_time": 2.5,
            },
            {
                "customer_id": "C2",
                "waiting_time": 1.5,
                "system_time": 4.0,
            },
        ]
    }


def test_csv_export_endpoint(client, payload):
    response = client.post(
        "/api/simulation/export/csv",
        json=payload,
    )

    assert response.status_code == 200
    assert response.mimetype == "text/csv"
    assert "attachment" in response.headers["Content-Disposition"]
    assert "customer_id" in response.get_data(as_text=True)
    assert "C1" in response.get_data(as_text=True)


def test_excel_export_endpoint(client, payload):
    response = client.post(
        "/api/simulation/export/excel",
        json=payload,
    )

    assert response.status_code == 200
    assert "attachment" in response.headers["Content-Disposition"]

    workbook = load_workbook(
        io.BytesIO(response.data)
    )
    worksheet = workbook.active

    assert worksheet.cell(1, 1).value == "customer_id"
    assert worksheet.cell(2, 1).value == "C1"


def test_csv_export_rejects_empty_rows(client):
    response = client.post(
        "/api/simulation/export/csv",
        json={"rows": []},
    )

    assert response.status_code == 400


def test_excel_export_rejects_empty_rows(client):
    response = client.post(
        "/api/simulation/export/excel",
        json={"rows": []},
    )

    assert response.status_code == 400


def test_csv_export_rejects_invalid_json_body(client):
    response = client.post(
        "/api/simulation/export/csv",
        data="invalid",
        content_type="text/plain",
    )

    assert response.status_code == 400

