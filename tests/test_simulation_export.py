import csv
import io
import json

import pytest
from openpyxl import load_workbook

from app.models.simulation_result import (
    CustomerSimulationResult,
    SimulationResult,
)
from app.services.export_service import ExportService


@pytest.fixture
def simulation_result():
    customer = CustomerSimulationResult(
        customer_id="C1",
        arrival_time=0.0,
        service_time=2.0,
        service_start=0.0,
        service_end=2.0,
        waiting_time=0.0,
        system_time=2.0,
        server_id=1,
        server_idle_time=0.0,
    )

    return SimulationResult(
        simulation_type="single_server",
        server_count=1,
        customers_processed=1,
        arrival_rate=0.5,
        service_rate=0.5,
        utilization=1.0,
        average_waiting_time=0.0,
        average_system_time=2.0,
        average_number_in_queue=0.0,
        average_number_in_system=1.0,
        throughput=0.5,
        total_idle_time=0.0,
        server_utilization={1: 1.0},
        customer_results=[customer],
    )


def test_simulation_to_csv_exports_customer_data(
    simulation_result,
):
    content = ExportService.simulation_to_csv(
        simulation_result
    )

    rows = list(csv.DictReader(io.StringIO(content)))

    assert len(rows) == 1
    assert rows[0]["customer_id"] == "C1"
    assert float(rows[0]["waiting_time"]) == 0.0


def test_simulation_to_csv_exports_summary(
    simulation_result,
):
    content = ExportService.simulation_to_csv(
        simulation_result,
        export_type="summary",
    )

    rows = list(csv.DictReader(io.StringIO(content)))

    assert len(rows) == 1
    assert rows[0]["customers_processed"] == "1"
    assert "average_waiting_time" in rows[0]


def test_simulation_to_excel_exports_customer_data(
    simulation_result,
):
    content = ExportService.simulation_to_excel(
        simulation_result
    )

    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook.active

    assert worksheet.title == "Customer Results"
    assert worksheet.cell(2, 1).value == "C1"


def test_simulation_to_excel_exports_summary(
    simulation_result,
):
    content = ExportService.simulation_to_excel(
        simulation_result,
        export_type="summary",
    )

    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook.active

    assert worksheet.title == "Simulation Summary"
    assert "average_waiting_time" in [
        cell.value for cell in worksheet[1]
    ]


def test_simulation_export_rejects_invalid_export_type(
    simulation_result,
):
    with pytest.raises(ValueError):
        ExportService.simulation_to_csv(
            simulation_result,
            export_type="invalid",
        )

