import csv
import io
import json

from app.services.result_formatter import ResultFormatter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class ExportService:
    """Export simulation data to CSV and Excel formats."""

    @staticmethod
    def _validate_rows(rows):
        if not isinstance(rows, list):
            raise ValueError("rows must be a list.")

        if not all(isinstance(row, dict) for row in rows):
            raise ValueError("Each row must be an object.")

        return rows

    @staticmethod
    def to_csv(rows):
        rows = ExportService._validate_rows(rows)

        if not rows:
            raise ValueError("Cannot export an empty dataset.")

        output = io.StringIO(newline="")
        fieldnames = list(
            dict.fromkeys(
                key for row in rows for key in row.keys()
            )
        )

        writer = csv.DictWriter(
            output,
            fieldnames=fieldnames,
            extrasaction="ignore",
        )

        writer.writeheader()
        writer.writerows(rows)

        return output.getvalue()

    @staticmethod
    def to_excel(rows, sheet_name="Simulation Results"):
        rows = ExportService._validate_rows(rows)

        if not rows:
            raise ValueError("Cannot export an empty dataset.")

        if not isinstance(sheet_name, str) or not sheet_name.strip():
            raise ValueError("sheet_name must be a non-empty string.")

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name[:31]

        fieldnames = list(
            dict.fromkeys(
                key for row in rows for key in row.keys()
            )
        )

        worksheet.append(fieldnames)

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

        for row in rows:
            excel_row = []

            for field in fieldnames:
                value = row.get(field)

                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(
                        value,
                        ensure_ascii=False,
                    )

                excel_row.append(value)

            worksheet.append(excel_row)



        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for column_index, field in enumerate(fieldnames, start=1):
            values = [
                str(worksheet.cell(row=row_index, column=column_index).value or "")
                for row_index in range(1, worksheet.max_row + 1)
            ]

            max_length = max(
                (len(value) for value in values),
                default=0,
            )

            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(max(max_length + 2, 12), 40)

        output = io.BytesIO()
        workbook.save(output)

        return output.getvalue()
    
    @staticmethod
    def simulation_to_csv(simulation_result, export_type="customers"):
        if export_type == "customers":
            rows = ResultFormatter.customer_rows(
                simulation_result
            )
        elif export_type == "summary":
            rows = ResultFormatter.summary_rows(
                simulation_result
            )
        else:
            raise ValueError(
                "export_type must be 'customers' or 'summary'."
            )

        return ExportService.to_csv(rows)

    @staticmethod
    def simulation_to_excel(
        simulation_result,
        export_type="customers",
    ):
        if export_type == "customers":
            rows = ResultFormatter.customer_rows(
                simulation_result
            )
            sheet_name = "Customer Results"
        elif export_type == "summary":
            rows = ResultFormatter.summary_rows(
                simulation_result
            )
            sheet_name = "Simulation Summary"
        else:
            raise ValueError(
                "export_type must be 'customers' or 'summary'."
            )

        return ExportService.to_excel(
            rows,
            sheet_name=sheet_name,
        )



